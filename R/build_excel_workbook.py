#!/usr/bin/env python3
"""
BUILD EXCEL WORKBOOK
Proyek: Riset Dinamika Inflasi ASEAN
Membuat file Excel 30 sheet 'ASEAN_Inflation_Quarterly_Final.xlsx' dan 'Data_Audit.xlsx'
dengan formatting profesional (Navy header, freeze panes, autofilter, conditional formatting, dll).
"""

import os
import glob
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print(">>> Memulai Pembangunan Workbook Excel Final (30 Sheet)...")

sheets_dir = "output/excel_sheets"
output_xlsx = "output/ASEAN_Inflation_Quarterly_Final.xlsx"
audit_xlsx = "output/Data_Audit.xlsx"

# Definisi Styles
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

font_data = Font(name="Arial", size=10, color="000000")
align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

thin_side = Side(border_style="thin", color="D9D9D9")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

fill_imputed = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
font_imputed = Font(name="Arial", size=10, color="B7791F", bold=True)

# Urutan 30 Sheet Wajib
sheet_order = [
    "README",
    "Data_Audit",
    "Raw_HCPI",
    "Raw_GDP",
    "Raw_Money_Growth",
    "Raw_REER",
    "Raw_FAO",
    "Raw_FFR",
    "Raw_Oil",
    "Quarterly_HCPI",
    "Quarterly_Inflation",
    "Quarterly_GDP",
    "Quarterly_Money",
    "Quarterly_REER",
    "Quarterly_FAO",
    "Quarterly_FFR",
    "Quarterly_Oil",
    "Panel_Quarterly_Raw",
    "Panel_Quarterly_Final",
    "Imputation_Audit",
    "Disaggregation_Audit",
    "Descriptive_Stats",
    "Stationarity",
    "Lag_Selection",
    "Model_Results",
    "Diagnostics",
    "IRF_Data",
    "FEVD_Data",
    "Model_Comparison",
    "Results_Summary"
]

wb = openpyxl.Workbook()
# Hapus default sheet
wb.remove(wb.active)

for sname in sheet_order:
    csv_path = os.path.join(sheets_dir, f"{sname}.csv")
    ws = wb.create_sheet(title=sname)
    
    if os.path.exists(csv_path):
        df = pd.read_csv(csv_path)
    else:
        df = pd.DataFrame({"Keterangan": ["Data belum tersedia untuk sheet ini"]})
    
    # Khusus sheet README: Buat tampilan dokumentasi berformat elegan
    if sname == "README":
        ws.views.sheetView[0].showGridLines = True
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 85
        
        # Title Banner
        ws.merge_cells("A1:B1")
        cell_t = ws["A1"]
        cell_t.value = "WORKBOOK RISET DINAMIKA INFLASI NEGARA-NEGARA ASEAN (2010–2025)"
        cell_t.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        cell_t.fill = PatternFill(start_color="1B365D", fill_type="solid")
        cell_t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35
        
        # Tulis baris README
        for r_idx, row in df.iterrows():
            row_num = r_idx + 3
            ws.row_dimensions[row_num].height = 22
            
            c_sec = ws.cell(row=row_num, column=1, value=str(row.iloc[0]) if pd.notna(row.iloc[0]) else "")
            c_desc = ws.cell(row=row_num, column=2, value=str(row.iloc[1]) if len(row) > 1 and pd.notna(row.iloc[1]) else "")
            
            is_header_row = str(row.iloc[0]).startswith("===") or str(row.iloc[0]).isupper()
            if is_header_row:
                c_sec.font = Font(name="Arial", size=11, bold=True, color="1B365D")
                c_desc.font = Font(name="Arial", size=11, bold=True, color="1B365D")
                c_sec.fill = PatternFill(start_color="E2E8F0", fill_type="solid")
                c_desc.fill = PatternFill(start_color="E2E8F0", fill_type="solid")
            else:
                c_sec.font = Font(name="Arial", size=10, bold=True, color="2D3748")
                c_desc.font = Font(name="Arial", size=10, color="1A202C")
                c_sec.border = border_cell
                c_desc.border = border_cell
                
        ws.freeze_panes = "A3"
        continue

    # Khusus sheet Results_Summary
    if sname == "Results_Summary":
        ws.views.sheetView[0].showGridLines = True
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 90
        
        ws.merge_cells("A1:B1")
        cell_t = ws["A1"]
        cell_t.value = "RINGKASAN EKSEKUTIF & TEMUAN UTAMA RISET INFLASI ASEAN"
        cell_t.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        cell_t.fill = PatternFill(start_color="1B365D", fill_type="solid")
        cell_t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35
        
        for r_idx, row in df.iterrows():
            row_num = r_idx + 3
            ws.row_dimensions[row_num].height = 22
            c_sec = ws.cell(row=row_num, column=1, value=str(row.iloc[0]) if pd.notna(row.iloc[0]) else "")
            c_desc = ws.cell(row=row_num, column=2, value=str(row.iloc[1]) if len(row) > 1 and pd.notna(row.iloc[1]) else "")
            
            if str(row.iloc[0]).startswith("==="):
                c_sec.font = Font(name="Arial", size=11, bold=True, color="1B365D")
                c_desc.font = Font(name="Arial", size=11, bold=True, color="1B365D")
                c_sec.fill = PatternFill(start_color="E2E8F0", fill_type="solid")
                c_desc.fill = PatternFill(start_color="E2E8F0", fill_type="solid")
            else:
                c_sec.font = Font(name="Arial", size=10, bold=True, color="2D3748")
                c_desc.font = Font(name="Arial", size=10, color="1A202C")
                c_sec.border = border_cell
                c_desc.border = border_cell
        ws.freeze_panes = "A3"
        continue

    # Untuk seluruh Data Tables lainnya
    ws.views.sheetView[0].showGridLines = True
    ws.row_dimensions[1].height = 26
    
    # Tulis Headers
    for c_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=c_idx, value=str(col_name))
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_cell

    # Tulis Rows Data
    flag_col_idx = None
    if "data_quality_flag" in df.columns:
        flag_col_idx = list(df.columns).index("data_quality_flag") + 1

    for r_idx, row in df.iterrows():
        row_num = r_idx + 2
        ws.row_dimensions[row_num].height = 19
        is_imputed_row = False
        if flag_col_idx and pd.notna(row.get("data_quality_flag")):
            if "Imputation" in str(row["data_quality_flag"]):
                is_imputed_row = True
                
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.border = border_cell
            
            if pd.isna(val):
                cell.value = ""
                cell.alignment = align_center
            else:
                # Tentukan tipe numerik atau teks
                if isinstance(val, (int, float)):
                    cell.value = val
                    cell.font = font_data
                    cell.alignment = align_right
                    col_str = str(df.columns[c_idx - 1]).lower()
                    if "year" in col_str or "quarter" in col_str or "n_" in col_str or "df" in col_str or "horizon" in col_str:
                        cell.number_format = "0"
                        cell.alignment = align_center
                    elif "growth" in col_str or "inflation" in col_str or "yoy" in col_str or "qoq" in col_str or "change" in col_str or "rate" in col_str or "pct" in col_str or "share" in col_str or "stat" in col_str or "estimate" in col_str or "r_squared" in col_str:
                        cell.number_format = "0.00"
                    elif "hcpi" in col_str or "reer" in col_str or "fao" in col_str or "oil" in col_str:
                        cell.number_format = "0.00"
                    elif "p_value" in col_str or "pval" in col_str:
                        cell.number_format = "0.0000"
                    else:
                        cell.number_format = "0.00"
                else:
                    val_str = str(val)
                    # Cek jika string numerik
                    try:
                        f_val = float(val_str)
                        cell.value = f_val
                        cell.number_format = "0.00"
                        cell.font = font_data
                        cell.alignment = align_right
                    except ValueError:
                        cell.value = val_str
                        cell.font = font_data
                        if len(val_str) <= 10:
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left
                            
            if is_imputed_row and str(df.columns[c_idx - 1]) in ["hcpi_q", "inflation_yoy", "money_growth", "reer_q"]:
                cell.fill = fill_imputed

    # Set Freeze Panes & AutoFilter
    last_col_letter = get_column_letter(len(df.columns))
    ws.auto_filter.ref = f"A1:{last_col_letter}{len(df) + 1}"
    
    if sname in ["Panel_Quarterly_Raw", "Panel_Quarterly_Final"]:
        ws.freeze_panes = "E2" # Freeze country, year, quarter, period + header
    elif "Quarterly_" in sname and "country" in df.columns:
        ws.freeze_panes = "B2"
    else:
        ws.freeze_panes = "A2"

    # Penyesuaian Lebar Kolom Otomatis
    for c_idx, col_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(c_idx)
        max_len = len(str(col_name))
        sample_vals = df[col_name].dropna().head(20).astype(str)
        if len(sample_vals) > 0:
            val_len = sample_vals.map(len).max()
            max_len = max(max_len, val_len)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 38)

# Simpan workbook utama
wb.save(output_xlsx)
print(f"Workbook final berhasil disimpan di: {output_xlsx}")

# ------------------------------------------------------------------------------
# BANGUN WORKBOOK DATA_AUDIT.XLSX
# ------------------------------------------------------------------------------
print(">>> Membangun Data_Audit.xlsx...")
wb_audit = openpyxl.Workbook()
wb_audit.remove(wb_audit.active)

audit_sheets = ["Data_Audit", "Audit_Sheet_Overview", "Audit_Global_Variables"]
for asname in audit_sheets:
    csv_p = os.path.join("output", f"{asname}.csv")
    if not os.path.exists(csv_p):
        csv_p = os.path.join(sheets_dir, f"{asname}.csv")
    
    if os.path.exists(csv_p):
        adf = pd.read_csv(csv_p)
    else:
        continue
        
    aws = wb_audit.create_sheet(title=asname.replace("Audit_", ""))
    aws.views.sheetView[0].showGridLines = True
    aws.row_dimensions[1].height = 26
    
    for c_idx, col_name in enumerate(adf.columns, 1):
        cell = aws.cell(row=1, column=c_idx, value=str(col_name))
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_cell
        
    for r_idx, row in adf.iterrows():
        row_num = r_idx + 2
        aws.row_dimensions[row_num].height = 19
        for c_idx, val in enumerate(row, 1):
            cell = aws.cell(row=row_num, column=c_idx)
            cell.border = border_cell
            cell.font = font_data
            if pd.isna(val):
                cell.value = ""
            elif isinstance(val, (int, float)):
                cell.value = val
                cell.alignment = align_right
                cell.number_format = "0.00" if isinstance(val, float) else "0"
            else:
                cell.value = str(val)
                cell.alignment = align_left
                
    last_col = get_column_letter(len(adf.columns))
    aws.auto_filter.ref = f"A1:{last_col}{len(adf) + 1}"
    aws.freeze_panes = "A2"
    
    for c_idx, col_name in enumerate(adf.columns, 1):
        col_letter = get_column_letter(c_idx)
        max_len = max(len(str(col_name)), adf[col_name].dropna().astype(str).map(len).max() if len(adf) > 0 else 10)
        aws.column_dimensions[col_letter].width = min(max_len + 4, 45)

wb_audit.save(audit_xlsx)
print(f"Data_Audit.xlsx berhasil disimpan di: {audit_xlsx}")
print(">>> Selesai Pembangunan Workbook Excel Final.")

